"""Validate the new designer's workflow boundary and actual editable output."""
import base64
import io
import os
import tempfile
import unittest
import zipfile
from xml.etree import ElementTree
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("NO_GCE_CHECK","true")

from mcp import Client
from pptx import Presentation
from pydantic import ValidationError
from backend.app.services.workspace_service import WorkspaceError, WorkspaceService
from backend.mcp_sastra.designer import DesignPlan, DesignerService
from backend.mcp_sastra.designer_demo import sample_source
from backend.mcp_sastra.sandbox import ROOT
from backend.mcp_sastra.server import create_server


class DesignerTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.service=WorkspaceService(Path(self.temp.name)/"designer.sqlite3")
        self.source=self.service.import_content(sample_source())["contents"][0]
        self.plan=DesignPlan.model_validate_json((ROOT/"examples/mcp/designed-lesson.json").read_text(encoding="utf-8"))
        self.designer=DesignerService(self.service)

    def tearDown(self):
        self.temp.cleanup()

    def test_invalid_asset_paths_and_repeated_content_are_rejected(self):
        raw=self.plan.model_dump()
        raw["slides"][1]["visual"]="../../secret.png"
        with self.assertRaises(ValidationError): DesignPlan.model_validate(raw)
        raw=self.plan.model_dump()
        raw["slides"][2]["items"][0]["body"]=raw["slides"][1]["items"][0]["body"]
        with self.assertRaises(ValidationError): DesignPlan.model_validate(raw)
        raw=self.plan.model_dump()
        raw["output"]="../../production.sqlite3"
        with self.assertRaises(ValidationError): DesignPlan.model_validate(raw)

    def test_approval_and_version_checked_before_renderer(self):
        with patch("backend.mcp_sastra.designer.run_renderer") as render:
            with self.assertRaises(WorkspaceError): self.designer.design(self.source["id"],1,self.plan)
            self.service.approve_content(self.source["id"],1)
            with self.assertRaises(WorkspaceError): self.designer.design(self.source["id"],1,self.plan)
            render.assert_not_called()

    def test_source_change_during_render_does_not_save_presentation(self):
        self.service.approve_content(self.source["id"],1)
        def racing_render(_):
            def change(w): w["contents"][0]["version"]+=1
            self.service.mutate(change)
            return b"fixture-not-a-pptx",{}
        with patch("backend.mcp_sastra.designer.run_renderer",side_effect=racing_render):
            with self.assertRaises(WorkspaceError): self.designer.design(self.source["id"],2,self.plan)
        db=self.designer._connect()
        self.assertEqual(db.execute("SELECT count(*) FROM presentations").fetchone()[0],0)
        db.close()

    async def test_actual_mcp_designer_and_pdf_resources(self):
        self.service.approve_content(self.source["id"],1)
        async with Client(create_server(self.service)) as client:
            result=await client.call_tool("design_presentation",{"assessment_id":self.source["id"],
                "expected_version":2,"plan":self.plan.model_dump()})
            self.assertFalse(result.is_error,result.content)
            metadata=result.structured_content
            data=base64.b64decode((await client.read_resource(metadata["resourceUri"])).contents[0].blob)
            deck=Presentation(io.BytesIO(data))
            with zipfile.ZipFile(io.BytesIO(data)) as package:
                types=ElementTree.fromstring(package.read("[Content_Types].xml"))
                self.assertTrue(all(not part.get("PartName") or part.get("PartName").lstrip("/") in package.namelist()
                    for part in types))
            self.assertEqual(len(deck.slides),8)
            self.assertAlmostEqual(deck.slide_width/deck.slide_height,16/9,places=5)
            text="\n".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
            self.assertIn("Rahasia",text)
            self.assertIn("Sumber energi untuk fotosintesis",text)
            self.assertNotIn(self.source["answerKey"][0]["rubric"],text)
            self.assertEqual(sum(shape.shape_type==13 for slide in deck.slides for shape in slide.shapes),3)
            preview=await client.call_tool("preview_presentation",{"presentation_id":metadata["id"]})
            self.assertFalse(preview.is_error,preview.content)
            pdf=base64.b64decode((await client.read_resource(preview.structured_content["resourceUri"])).contents[0].blob)
            self.assertTrue(pdf.startswith(b"%PDF-"))


if __name__=="__main__": unittest.main()
