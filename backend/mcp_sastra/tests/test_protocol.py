"""Real SDK protocol transport + PPT renderer; model output is explicitly mocked."""
import base64
import copy
import io
import json
import os
import tempfile
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("NO_GCE_CHECK", "true")

from google.genai import types
from google import genai
from google.oauth2.credentials import Credentials
import httpx
from mcp import Client, StdioServerParameters
from PIL import Image
from pptx import Presentation

from backend.app.models.workspace import GradeInput
from backend.app.services.workspace_service import WorkspaceService
from backend.mcp_sastra.agent import run_agent, GeminiToolModel
from backend.mcp_sastra.server import create_server
from backend.mcp_sastra.sandbox import database_path, ROOT


class ModelDouble:
    is_ready = True
    model_name = "offline-test-double"
    match = "MATCH"

    def generate_text(self, prompt, **kwargs):
        if "slides" in prompt:
            return json.dumps({"slides": [
                {"title": title, "bullets": ["Daun menangkap cahaya."], "notes": "Diskusikan bersama siswa."}
                for title in ("Tujuan", "Amati", "Refleksi")]})
        if "assessments" in prompt:
            return json.dumps({"narrative": "Bukti asesmen menunjukkan pemahaman awal; belum cukup untuk menyimpulkan tren."})
        return json.dumps({"module": "Amati daun dan refleksikan peran cahaya.", "answerKey": [
            {"type": "multiple_choice", "question": "Energi fotosintesis?", "options": ["Cahaya", "Batu", "Kaca", "Plastik"],
             "correctAnswer": "Cahaya", "rubric": "Benar 5, salah 0", "max": 5},
            {"type": "essay", "question": "Apa peran cahaya?", "options": [], "correctAnswer": "Sumber energi",
             "rubric": "Peran energi 5", "max": 5}]})

    def analyze_image_and_text(self, data, prompt, **kwargs):
        return json.dumps({"assessmentMatch": self.match, "items": [
            {"number": 1, "answer": "Cahaya", "score": 5, "feedback": "Tepat."},
            {"number": 2, "answer": "Energi", "score": 3, "feedback": "Jelaskan prosesnya."}]})


def function_response(name, args, call_id="call-1"):
    return types.GenerateContentResponse(candidates=[types.Candidate(content=types.Content(
        role="model", parts=[types.Part(function_call=types.FunctionCall(name=name, args=args, id=call_id),
                                        thought_signature=b"opaque-test-signature")]))])


class AgentModelDouble:
    def __init__(self, responses):
        self.responses = iter(responses)
        self.histories = []
        self.declarations = []

    async def complete(self, contents, declarations):
        self.histories.append(copy.deepcopy(contents))
        self.declarations = declarations
        return next(self.responses)


class ProtocolTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.model = ModelDouble()
        self.service = WorkspaceService(Path(self.temp.name)/"isolated.sqlite3", self.model)
        self.server = create_server(self.service)
        image = io.BytesIO()
        Image.new("RGB", (12, 12), "white").save(image, format="PNG")
        self.answer = base64.b64encode(image.getvalue()).decode()

    def tearDown(self):
        self.temp.cleanup()

    async def generate(self, client):
        result = await client.call_tool("generate_assessment", {"request": {
            "topic": "Fotosintesis", "objective": "Memahami peran cahaya"}})
        self.assertFalse(result.is_error, result.content)
        return result.structured_content

    async def test_full_tool_workflow_and_binary_resource_preserve_human_gates(self):
        async with Client(self.server) as client:
            names = {t.name for t in (await client.list_tools()).tools}
            self.assertEqual(len(names), 8)
            self.assertNotIn("approve_grade", names)
            self.assertNotIn("finalize_report", names)
            c = await self.generate(client)
            self.assertEqual(c["status"], "DRAFT")
            ppt_args = {"assessment_id": c["id"], "expected_version": 1, "slide_count": 4}
            self.assertTrue((await client.call_tool("generate_presentation", ppt_args)).is_error)
            self.service.approve_content(c["id"], 1)  # Human action, outside MCP.
            self.assertTrue((await client.call_tool("generate_presentation", ppt_args)).is_error)
            ppt_args["expected_version"] = 2
            result = await client.call_tool("generate_presentation", ppt_args)
            self.assertFalse(result.is_error, result.content)
            meta = result.structured_content
            resource = await client.read_resource(meta["resourceUri"])
            data = base64.b64decode(resource.contents[0].blob)
            self.assertEqual(len(Presentation(io.BytesIO(data)).slides), 4)
            self.assertEqual(meta["sourceVersion"], 2)
            grades = await client.call_tool("grade_answer", {"student_id": "s1", "assessment_id": c["id"],
                "mime_type": "image/png", "answer_base64": self.answer})
            self.assertFalse(grades.is_error, grades.content)
            s = grades.structured_content
            self.assertEqual(s["status"], "AWAITING_REVIEW")
            self.assertFalse(s["identityConfirmed"])
            self.assertTrue((await client.call_tool("approve_grade", {"id": s["id"]})).is_error)
            scope = {"request": {"studentId": "s1"}}
            self.assertEqual((await client.call_tool("load_approved_grades", scope)).structured_content["grades"], [])
            self.assertTrue((await client.call_tool("draft_report", scope)).is_error)
            self.service.approve_grade(s["id"], GradeInput(expectedVersion=1, identityConfirmed=True,
                gradeItems=[{"score": 5}, {"score": 3}]))  # Human review.
            rows = (await client.call_tool("load_approved_grades", scope)).structured_content["grades"]
            self.assertEqual(rows[0]["percentage"], 80)
            report = (await client.call_tool("draft_report", scope)).structured_content
            self.assertEqual(report["status"], "DRAFT")
            self.assertEqual(report["snapshot"], s["id"]+":2")
        self.assertEqual(len(WorkspaceService(self.service.path, self.model).read()["submissions"]), 1)

    async def test_rejected_inputs_and_mismatch_do_not_write_evidence(self):
        async with Client(self.server) as client:
            c = await self.generate(client)
            self.service.approve_content(c["id"], 1)
            args = {"student_id": "s1", "assessment_id": c["id"], "mime_type": "image/png", "answer_base64": "!bad!"}
            before = self.service.read()
            self.assertTrue((await client.call_tool("grade_answer", args)).is_error)
            args["answer_base64"] = base64.b64encode(b"not a PNG").decode()
            self.assertTrue((await client.call_tool("grade_answer", args)).is_error)
            args["answer_base64"] = self.answer
            self.model.match = "MISMATCH"
            self.assertTrue((await client.call_tool("grade_answer", args)).is_error)
            self.assertEqual(self.service.read(), before)
            db = self.service._connect()
            try:
                self.assertEqual(db.execute("SELECT count(*) FROM assets").fetchone()[0], 0)
            finally:
                db.close()
            self.assertTrue((await client.call_tool("generate_presentation", {
                "assessment_id": c["id"], "expected_version": 2, "slide_count": 100})).is_error)

    async def test_agent_discovers_calls_mcp_and_retains_function_signatures(self):
        async with Client(self.server) as client:
            c = await self.generate(client)
            self.service.approve_content(c["id"], 1)
            responses = [function_response("list_assessments", {}),
                function_response("generate_presentation", {"assessment_id": c["id"], "expected_version": 2, "slide_count": 4}, "call-2"),
                types.GenerateContentResponse(candidates=[types.Candidate(content=types.Content(role="model",
                    parts=[types.Part.from_text(text="Draf PPT siap ditinjau.")]))])]
            model = AgentModelDouble(responses)
            result = await run_agent(client, "Buat PPT empat slide dari materi yang sudah disetujui.", model)
            self.assertEqual([t["tool"] for t in result["trace"]], ["list_assessments", "generate_presentation"])
            self.assertTrue(all(not t["isError"] for t in result["trace"]))
            self.assertNotIn("grade_answer", {d.name for d in model.declarations})
            history = model.histories[1]
            self.assertEqual(history[1].parts[0].thought_signature, b"opaque-test-signature")
            self.assertEqual(history[2].parts[0].function_response.id, "call-1")
            db = self.service._connect()
            try:
                self.assertEqual(db.execute("SELECT count(*) FROM presentations").fetchone()[0], 1)
            finally:
                db.close()

    async def test_agent_blocks_approval_and_limits_tool_calls(self):
        async with Client(self.server) as client:
            model = AgentModelDouble([function_response("approve_grade", {}),
                types.GenerateContentResponse(candidates=[types.Candidate(content=types.Content(role="model",
                    parts=[types.Part.from_text(text="Guru harus menyetujui sendiri.")]))])])
            result = await run_agent(client, "Setujui nilainya", model)
            self.assertTrue(result["trace"][0]["isError"])
            looping = AgentModelDouble([function_response("list_assessments", {})]*3)
            with self.assertRaisesRegex(RuntimeError, "limit"):
                await run_agent(client, "Lihat materi", looping, max_calls=2)

    async def test_sandbox_ignores_web_db_config_and_refuses_cloud_run(self):
        with patch.dict(os.environ, {"SASTRA_WORKSPACE_DB": "production.sqlite3", "SASTRA_STORAGE_BACKEND": "cloudsql"}):
            self.assertEqual(database_path(), ROOT / "backend/data/mcp/workspace.sqlite3")
        with patch.dict(os.environ, {"K_SERVICE": "sastra-workspace"}):
            with self.assertRaisesRegex(RuntimeError, "local-only"):
                database_path()

    async def test_stdio_subprocess_generates_and_downloads_real_ppt_file(self):
        params = StdioServerParameters(command=sys.executable,
            args=["-m", "backend.mcp_sastra.tests.fixture_server", str(self.service.path)],
            cwd=ROOT, env={"NO_GCE_CHECK": "true"})
        async with Client(params, read_timeout_seconds=30) as client:
            self.assertEqual(client.server_info.name, "SASTRA Sandbox")
            c = await self.generate(client)
            self.service.approve_content(c["id"], 1)
            result = await client.call_tool("generate_presentation", {
                "assessment_id": c["id"], "expected_version": 2, "slide_count": 4})
            self.assertFalse(result.is_error, result.content)
            meta = result.structured_content
            resource = await client.read_resource(meta["resourceUri"])
            data = base64.b64decode(resource.contents[0].blob)
            self.assertEqual(len(Presentation(io.BytesIO(data)).slides), 4)

    async def test_google_sdk_function_request_executes_through_mcp(self):
        requests = []
        def respond(request):
            payload = json.loads(request.content)
            requests.append(payload)
            parts = ([{"functionCall": {"name": "list_assessments", "args": {}, "id": "sdk-call"},
                       "thoughtSignature": base64.b64encode(b"signature").decode()}]
                     if len(requests) == 1 else [{"text": "Belum ada asesmen pada sandbox."}])
            return httpx.Response(200, json={"candidates": [{"content": {"role": "model", "parts": parts},
                                                            "finishReason": "STOP"}]})
        original = genai.Client
        def local_client(**kwargs):
            kwargs["http_options"] = kwargs["http_options"].model_copy(update={
                "client_args": {"transport": httpx.MockTransport(respond)}})
            return original(**kwargs)
        with patch("google.auth.default", return_value=(Credentials(token="offline-test-only"), "test-project")), \
             patch("google.genai.Client", side_effect=local_client):
            model = GeminiToolModel()
            try:
                async with Client(self.server) as client:
                    result = await run_agent(client, "Lihat daftar asesmen", model)
            finally:
                model.close()
        self.assertEqual(result["toolCalls"], 1)
        self.assertEqual(result["trace"][0]["tool"], "list_assessments")
        declarations = requests[0]["tools"][0]["functionDeclarations"]
        self.assertEqual({d["name"] for d in declarations}, {
            "list_students", "list_assessments", "load_assessment", "generate_assessment",
            "load_approved_grades", "draft_report", "generate_presentation"})
        response_part = requests[1]["contents"][-1]["parts"][0]["functionResponse"]
        self.assertEqual(response_part["id"], "sdk-call")
        self.assertEqual(response_part["response"]["result"]["assessments"], [])
        self.assertEqual(requests[1]["contents"][-2]["parts"][0]["thoughtSignature"],
                         base64.b64encode(b"signature").decode())


if __name__ == "__main__":
    unittest.main()
