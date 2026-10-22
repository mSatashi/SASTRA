import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from pptx import Presentation

from backend.app.main import app
from backend.app.api import workspace_routes
from backend.app.config import settings
from backend.app.services.workspace_service import WorkspaceService
from backend.app.services.presentation_service import PresentationService


class SlideModel:
    is_ready = True
    model_name = "slide-contract-stub"

    def generate_text(self, prompt, **kwargs):
        self.prompt = prompt
        plan=json.loads((Path(__file__).resolve().parents[2]/"examples/mcp/designed-lesson.json").read_text(encoding="utf-8"))
        plan.update(title="Fotosintesis",coverVisual="none")
        plan["slides"]=[plan["slides"][i] for i in (0,1,6)]
        return json.dumps(plan)


class PresentationContractTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.model = SlideModel()
        self.service = WorkspaceService(Path(self.temp.name) / "test.sqlite3", self.model)
        def seed(w):
            w["contents"] = [{"id": "c1", "title": "Fotosintesis", "status": "APPROVED", "version": 2,
                "objective": "Menjelaskan peran cahaya.", "module": "Daun menggunakan klorofil untuk menangkap cahaya.",
                "subject": "IPAS", "className": "IV A", "questions": ["Apa peran cahaya?"],
                "answerKey": [{"correctAnswer": "PRIVATE_KEY_SENTINEL"}], "rubric": "PRIVATE_RUBRIC_SENTINEL"}]
        self.service.mutate(seed)
        self.route_patch = patch.object(workspace_routes, "workspace_service", self.service)
        self.route_patch.start()
        self.client = TestClient(app)
        self.url = "/api/v1/assessment-versions/c1/presentations"

    def tearDown(self):
        self.client.close()
        self.route_patch.stop()
        self.temp.cleanup()

    def generate(self, **kwargs):
        return self.client.post(self.url, json={"expectedVersion": 2, "slideCount": 4, **kwargs})

    def test_editable_file_and_metadata_survive_new_service_connection(self):
        response = self.generate()
        self.assertEqual(response.status_code, 200, response.text)
        draft = response.json()
        self.assertEqual(draft["sourceVersion"], 2)
        self.assertEqual(draft["slideCount"], 4)
        self.assertNotIn("PRIVATE_KEY_SENTINEL", self.model.prompt)
        self.assertNotIn("PRIVATE_RUBRIC_SENTINEL", self.model.prompt)
        file = self.client.get(f'/api/v1/presentations/{draft["id"]}/file')
        self.assertEqual(file.status_code, 200)
        self.assertIn("attachment", file.headers["content-disposition"])
        prs = Presentation(io.BytesIO(file.content))
        self.assertEqual(len(prs.slides), 4)
        self.assertAlmostEqual(prs.slide_width / prs.slide_height, 16 / 9, places=5)
        self.assertIn("Fotosintesis", prs.slides[0].shapes[1].text)
        self.assertIn("materi c1 versi 2", prs.slides[1].notes_slide.notes_text_frame.text)
        reopened = PresentationService(WorkspaceService(self.service.path, self.model))
        self.assertEqual(reopened.download(draft["id"]), file.content)
        self.assertEqual(reopened.list_for("c1")[0]["id"], draft["id"])

    def test_rejects_unapproved_stale_invalid_and_unavailable_inputs(self):
        self.assertEqual(self.generate(expectedVersion=1).status_code, 409)
        for count in (3, 13, True):
            self.assertEqual(self.generate(slideCount=count).status_code, 422)
        self.model.is_ready = False
        self.assertEqual(self.generate().status_code, 503)
        self.model.is_ready = True
        self.service.mutate(lambda w: w["contents"][0].update(status="DRAFT"))
        self.assertEqual(self.generate().status_code, 422)
        self.assertEqual(self.client.get('/api/v1/presentations/missing/file').status_code, 404)

    def test_invalid_model_result_and_source_race_do_not_save_file(self):
        for raw in ('{"slides":[]}', '{"slides":'+json.dumps([{"title":"T","bullets":["x"*111]}]*3)+'}'):
            with patch.object(self.model, "generate_text", return_value=raw):
                self.assertEqual(self.generate().status_code, 502)
        with patch.object(self.model, "generate_text", return_value=SlideModel().generate_text("")):
            self.assertEqual(self.generate(slideCount=6).status_code, 502)
        def racing(prompt, **kwargs):
            self.service.mutate(lambda w: w["contents"][0].update(version=3))
            return SlideModel().generate_text(prompt)
        with patch.object(self.model, "generate_text", side_effect=racing):
            self.assertEqual(self.generate().status_code, 409)
        self.assertEqual(self.client.get(self.url).json(), [])

    def test_generation_and_download_remain_behind_team_session(self):
        with patch.object(settings, "WEB_PUBLIC_MODE", True), patch.object(settings, "WEB_ACCESS_CODE", "test-code-with-at-least-24-characters"):
            self.assertEqual(self.generate().status_code, 401)
            self.assertEqual(self.client.get('/api/v1/presentations/missing/file').status_code, 401)
